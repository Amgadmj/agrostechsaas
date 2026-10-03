"""
Convert Agrostech Novel R&D Markdown Documents to Formatted Word (.docx) Files.
Handles headings, styled tables, bullet/numbered lists, code/box blocks, and embedded charts.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DIR_PATH = r"c:\Users\Razer\Desktop\Agrostech\novel_rd_directions"

# Styling Palette (Institutional AgTech Theme: Navy / Deep Teal / Steel)
PRIMARY_COLOR = RGBColor(16, 44, 87)       # Navy
SECONDARY_COLOR = RGBColor(28, 110, 164)   # Blue
TEXT_COLOR = RGBColor(33, 37, 41)          # Dark Slate
MUTED_COLOR = RGBColor(108, 117, 125)      # Gray
CODE_BG = "F8F9FA"
TABLE_HEADER_BG = "102C57"
TABLE_ALT_BG = "F4F6F9"
BORDER_COLOR = "D0D5DD"


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_formatted_runs(paragraph, text, base_font="Calibri", base_size=11, base_color=TEXT_COLOR):
    """Parses inline bold, italic, and inline code markdown."""
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = base_font
        run.font.size = Pt(base_size)
        run.font.color.rgb = base_color

        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run.text = part[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            run.text = part


def process_markdown_file(md_path, docx_path, attach_images=False):
    print(f"[*] Processing {os.path.basename(md_path)} -> {os.path.basename(docx_path)}...")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Document Geometry
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal code_block_lines
        if not code_block_lines:
            return
        code_text = "".join(code_block_lines).strip("\n")
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, CODE_BG)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        code_block_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        parsed_rows = []
        for tl in table_lines:
            tl = tl.strip()
            if not tl or re.match(r'^\|?\s*[-:]+[-| :]*$', tl):
                continue
            cols = [c.strip() for c in tl.split('|')]
            if cols and cols[0] == '':
                cols = cols[1:]
            if cols and cols[-1] == '':
                cols = cols[:-1]
            if cols:
                parsed_rows.append(cols)

        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row in enumerate(parsed_rows):
                is_header = (r_idx == 0)
                for c_idx in range(num_cols):
                    cell = table.cell(r_idx, c_idx)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    val = row[c_idx] if c_idx < len(row) else ""
                    if is_header:
                        set_cell_background(cell, TABLE_HEADER_BG)
                        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        run = p.add_run(val)
                        run.font.name = "Calibri"
                        run.font.size = Pt(9.5)
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        bg = TABLE_ALT_BG if (r_idx % 2 == 1) else "FFFFFF"
                        set_cell_background(cell, bg)
                        set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        add_formatted_runs(p, val, base_size=9)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
        table_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle Code Fences
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code_block()
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
            continue

        if in_code_block:
            code_block_lines.append(line)
            continue

        # Handle Markdown Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            in_table = False
            flush_table()

        if not stripped:
            continue

        # Horizontal Rule
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="{BORDER_COLOR}"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            continue

        # Headings
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[2:])
            run.font.name = "Calibri"
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = PRIMARY_COLOR
            continue
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(stripped[3:])
            run.font.name = "Calibri"
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = SECONDARY_COLOR
            continue
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[4:])
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = PRIMARY_COLOR
            continue
        elif stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[5:])
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = SECONDARY_COLOR
            continue

        # Blockquote / Alerts
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="24" w:space="0" w:color="102C57"/>'
                f'<w:top w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'<w:bottom w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, text, base_size=10, base_color=RGBColor(51, 65, 85))
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # Bullet List Items
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, stripped[2:])
            continue

        # Numbered List Items
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, num_match.group(2))
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        add_formatted_runs(p, stripped)

    # Flush remaining state
    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    # Append Visual Exhibits if requested (e.g. for the Executive Dossier)
    if attach_images:
        print("[*] Embedding benchmark visual charts into document...")
        doc.add_page_break()
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        run = h.add_run("Appendix: Empirical Benchmark Visual Exhibits & VC Dashboard")
        run.font.name = "Calibri"
        run.font.size = Pt(18)
        run.bold = True
        run.font.color.rgb = PRIMARY_COLOR

        charts = [
            ("Figure 1: Unified Multi-Panel Due Diligence Dashboard (All 3 Candidates)", os.path.join(DIR_PATH, "benchmark_summary_dashboard.png")),
            ("Figure 2: Candidate 1 — Área Contínua Multi-Feature Synchrony & ROC Curves", os.path.join(DIR_PATH, "benchmark_area_continua.png")),
            ("Figure 3: Candidate 2 — Value-of-Information (VoI) Drone Dispatch & Economic Payoff", os.path.join(DIR_PATH, "benchmark_voi_drone.png")),
            ("Figure 4: Candidate 3 — Uncertainty-Aware SCS-CN Monte Carlo & Terrace Breach Risk", os.path.join(DIR_PATH, "benchmark_runoff_erosion.png")),
        ]

        for caption, img_path in charts:
            if os.path.exists(img_path):
                cp = doc.add_paragraph()
                cp.paragraph_format.space_before = Pt(12)
                cp.paragraph_format.space_after = Pt(4)
                c_run = cp.add_run(caption)
                c_run.font.name = "Calibri"
                c_run.font.size = Pt(12)
                c_run.bold = True
                c_run.font.color.rgb = SECONDARY_COLOR

                ip = doc.add_paragraph()
                ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ip.paragraph_format.space_before = Pt(2)
                ip.paragraph_format.space_after = Pt(10)
                ip.add_run().add_picture(img_path, width=Inches(6.8))

    doc.save(docx_path)
    print(f"[SUCCESS] Saved: {docx_path} ({os.path.getsize(docx_path)} bytes)")


def main():
    docs_to_convert = [
        {
            "md": os.path.join(DIR_PATH, "AGROSTECH_NOVEL_RD_DOSSIER.md"),
            "docx": os.path.join(DIR_PATH, "AGROSTECH_NOVEL_RD_DOSSIER.docx"),
            "attach_images": True
        },
        {
            "md": os.path.join(DIR_PATH, "MATHEMATICAL_FORMULATIONS_AND_ALGORITHMS.md"),
            "docx": os.path.join(DIR_PATH, "MATHEMATICAL_FORMULATIONS_AND_ALGORITHMS.docx"),
            "attach_images": False
        },
        {
            "md": os.path.join(DIR_PATH, "PRIOR_ART_AND_PATENT_CLEARANCE.md"),
            "docx": os.path.join(DIR_PATH, "PRIOR_ART_AND_PATENT_CLEARANCE.docx"),
            "attach_images": False
        },
    ]

    for item in docs_to_convert:
        if os.path.exists(item["md"]):
            process_markdown_file(item["md"], item["docx"], attach_images=item["attach_images"])
        else:
            print(f"[WARN] File not found: {item['md']}")


if __name__ == "__main__":
    main()
